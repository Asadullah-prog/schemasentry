"""Dataset readers for CSV, Excel (.xlsx), JSON arrays, and JSON Lines.

All reading is local. Values are returned as strings (or ``None`` for missing);
type interpretation happens later in the validation engine, which treats a
string like ``"42"`` as a valid ``integer``. This keeps CSV and Excel behavior
consistent.
"""

from __future__ import annotations

import csv
import json
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import openpyxl


class DatasetError(Exception):
    """Raised when a dataset file cannot be read or is structurally invalid."""


@dataclass
class Dataset:
    """Tabular data read from a local file."""

    columns: list[str]
    rows: list[dict[str, str | None]]
    source: str
    format: str  # "csv" | "xlsx" | "json" | "jsonl"
    sheet: str | None = None
    row_count: int = field(init=False)

    def __post_init__(self) -> None:
        self.row_count = len(self.rows)


def _check_headers(columns: list[str], source: str) -> None:
    if not columns:
        raise DatasetError(f"{source}: no header row found; the dataset is empty.")
    seen: set[str] = set()
    for name in columns:
        if name is None or str(name).strip() == "":
            raise DatasetError(f"{source}: blank column name in header row.")
        if name in seen:
            raise DatasetError(f"{source}: duplicate column name {name!r} in header row.")
        seen.add(name)


def read_csv(path: Path) -> Dataset:
    """Read a CSV file (UTF-8, with or without BOM; quoted fields supported)."""
    source = str(path)
    try:
        text_stream = path.open("r", encoding="utf-8-sig", newline="")
    except OSError as exc:
        raise DatasetError(f"{source}: cannot open file: {exc.strerror or exc}") from exc
    with text_stream:
        try:
            peek = text_stream.read(1)
        except UnicodeDecodeError as exc:
            raise DatasetError(f"{source}: file is not valid UTF-8 text.") from exc
        if peek == "":
            raise DatasetError(f"{source}: file is empty.")
        text_stream.seek(0)
        try:
            reader = csv.DictReader(text_stream)
        except csv.Error as exc:
            raise DatasetError(f"{source}: malformed CSV: {exc}") from exc
        columns = list(reader.fieldnames or [])
        _check_headers(columns, source)
        rows: list[dict[str, str | None]] = []
        try:
            for lineno, record in enumerate(reader, start=2):
                if record is None:
                    continue
                extras = record.pop(None, None)
                if extras:
                    raise DatasetError(
                        f"{source}: row {lineno} has more fields than the header row."
                    )
                # Skip fully-blank rows (trailing newlines).
                if all(v is None or v == "" for v in record.values()):
                    continue
                rows.append(
                    {
                        col: (None if (val is None or val == "") else val)
                        for col, val in record.items()
                    }
                )
        except csv.Error as exc:
            raise DatasetError(f"{source}: malformed CSV: {exc}") from exc
    return Dataset(columns=columns, rows=rows, source=source, format="csv")


def read_xlsx(path: Path, sheet: str | None = None) -> Dataset:
    """Read a worksheet from an .xlsx workbook.

    Uses the named sheet when ``sheet`` is given, otherwise the active sheet.
    """
    source = str(path)
    try:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises several zip/parse error types
        raise DatasetError(f"{source}: cannot open Excel workbook: {exc}") from exc
    try:
        if sheet is not None:
            if sheet not in workbook.sheetnames:
                raise DatasetError(
                    f"{source}: sheet {sheet!r} not found "
                    f"(available: {', '.join(workbook.sheetnames) or 'none'})."
                )
            worksheet = workbook[sheet]
        else:
            worksheet = workbook.active
            sheet = worksheet.title
        raw_rows = list(worksheet.iter_rows(values_only=True))
    finally:
        workbook.close()
    if not raw_rows:
        raise DatasetError(f"{source}: sheet {sheet!r} is empty.")
    header = ["" if v is None else str(v) for v in raw_rows[0]]
    _check_headers(header, source)
    columns = header
    rows: list[dict[str, str | None]] = []
    for lineno, raw in enumerate(raw_rows[1:], start=2):
        values = list(raw) + [None] * (len(columns) - len(raw))
        if len(raw) > len(columns):
            raise DatasetError(f"{source}: row {lineno} has more fields than the header row.")
        record = {
            col: (None if (val is None or (isinstance(val, str) and val == "")) else str(val))
            for col, val in zip(columns, values, strict=True)
        }
        if all(v is None for v in record.values()):
            continue
        rows.append(record)
    return Dataset(columns=columns, rows=rows, source=source, format="xlsx", sheet=sheet)


def _rows_from_objects(objects: list[Any], source: str) -> Dataset:
    if not objects:
        raise DatasetError(f"{source}: JSON dataset contains no records.")
    columns: list[str] = []
    for obj in objects:
        if not isinstance(obj, dict):
            raise DatasetError(f"{source}: expected an array of objects.")
        for key in obj:
            if key not in columns:
                columns.append(str(key))
    _check_headers(columns, source)
    rows = [
        {
            col: (None if (obj.get(col) is None or obj.get(col) == "") else str(obj.get(col)))
            for col in columns
        }
        for obj in objects
    ]
    return Dataset(columns=columns, rows=rows, source=source, format="json")


def read_json(path: Path) -> Dataset:
    """Read a JSON array of objects, falling back to JSON Lines."""
    source = str(path)
    try:
        text = path.read_text(encoding="utf-8-sig")
    except OSError as exc:
        raise DatasetError(f"{source}: cannot open file: {exc.strerror or exc}") from exc
    except UnicodeDecodeError as exc:
        raise DatasetError(f"{source}: file is not valid UTF-8 text.") from exc
    if not text.strip():
        raise DatasetError(f"{source}: file is empty.")
    try:
        parsed = json.loads(text)
    except json.JSONDecodeError:
        parsed = None
    if isinstance(parsed, list):
        return _rows_from_objects(parsed, source)
    if isinstance(parsed, dict):
        raise DatasetError(f"{source}: expected a JSON array of objects or JSON Lines.")
    # Fall back to JSON Lines.
    objects: list[Any] = []
    for lineno, line in enumerate(text.splitlines(), start=1):
        if not line.strip():
            continue
        try:
            objects.append(json.loads(line))
        except json.JSONDecodeError as exc:
            raise DatasetError(f"{source}: line {lineno} is not valid JSON.") from exc
    if not objects:
        raise DatasetError(f"{source}: file contains no records.")
    dataset = _rows_from_objects(objects, source)
    dataset.format = "jsonl"
    return dataset


_FORMAT_READERS = {
    ".csv": read_csv,
    ".xlsx": read_xlsx,
    ".json": read_json,
    ".jsonl": read_json,
    ".ndjson": read_json,
}


def read_dataset(path: str | Path, sheet: str | None = None, format: str | None = None) -> Dataset:
    """Read a dataset file, inferring the format from its extension.

    ``format`` overrides inference; one of ``csv``, ``xlsx``, ``json``,
    ``jsonl``. ``sheet`` selects an Excel worksheet by name.
    """
    file_path = Path(path)
    if not file_path.is_file():
        raise DatasetError(f"{path}: file not found.")
    fmt = (format or file_path.suffix).lower().lstrip(".")
    if fmt not in ("csv", "xlsx", "json", "jsonl", "ndjson"):
        raise DatasetError(
            f"{path}: unsupported format {fmt!r}; "
            "supported extensions are .csv, .xlsx, .json, .jsonl, .ndjson."
        )
    if sheet is not None and fmt != "xlsx":
        raise DatasetError(f"{path}: --sheet only applies to Excel (.xlsx) files.")
    if fmt == "csv":
        return read_csv(file_path)
    if fmt == "xlsx":
        return read_xlsx(file_path, sheet=sheet)
    return read_json(file_path)
